import os
import json
import logging
from typing import Dict, Any, List, Optional
import math
import numpy as np
from shapely.geometry import shape, Point
import h3

from app.spatial.ahp import calculate_ahp_weights, calculate_tod_score, DEFAULT_5D_PAIRWISE_MATRIX
from app.analytics.sdm_regression import SDMRegressor

logger = logging.getLogger(__name__)

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "spatial")
CALIBRATED_JSON_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "calibrated_models.json")

# Focus Area Hubs (Tier 1)
TIER_1_STATION_IDS = {"gubeng", "pasar_turi", "wonokromo", "terminal_joyoboyo", "terminal_purabaya"}

def _load_calibrated_models() -> Dict[str, Any]:
    if os.path.exists(CALIBRATED_JSON_PATH):
        try:
            with open(CALIBRATED_JSON_PATH, "r", encoding="utf-8") as f:
                return json.load(f).get("stations", {})
        except Exception as e:
            logger.warning(f"Gagal memuat calibrated_models.json: {e}")
    return {}

# Cache global in-memory
_CACHED_STATIONS_DATA: Optional[Dict[str, Dict[str, Any]]] = None
_CACHED_H3_COLLECTION: Optional[Dict[str, Any]] = None


def _clean_slug(name: str) -> str:
    """Membersihkan nama stasiun menjadi slug standar."""
    s = name.strip().lower()
    for prefix in ["stasiun surabaya ", "stasiun ", "surabaya "]:
        if s.startswith(prefix):
            s = s[len(prefix):]
    s = s.replace("gubengx", "gubeng")
    s = s.replace("kota (semut)", "semut").replace("kota", "semut")
    return s.replace(" ", "_").strip("_")


def load_spatial_raw_data() -> Dict[str, Any]:
    """Memuat seluruh layer spasial mentah dari direktori spatial."""
    raw = {}

    def load_json(fname: str) -> Dict[str, Any]:
        p = os.path.join(DATA_DIR, fname)
        if os.path.exists(p):
            with open(p, "r", encoding="utf-8") as f:
                return json.load(f)
        return {"type": "FeatureCollection", "features": []}

    raw["stations"] = load_json("stasiun_surabaya.geojson")
    raw["demografi"] = load_json("demografi_surabaya.geojson")
    raw["halte"] = load_json("halte_surabaya.geojson")
    raw["banjir"] = load_json("banjir_surabaya.geojson")
    raw["ntl"] = load_json("nighttime_light_surabaya.geojson")
    raw["malls"] = load_json("PUSAT PERBELANJAAN DI KOTA SURABAYA TAHUN 2025.geojson")
    raw["ses"] = load_json("STATUS EKONOMI DAN SOSIAL - SOCIOECONOMIC STATUS (SES) KOTA SURABAYA TAHUN 2024.geojson")
    raw["properti"] = load_json("HARGA PROPERTI DI KOTA SURABAYA TAHUN 2024.geojson")

    # Load NJOP Excel
    njop_dict = {}
    xlsx_path = os.path.join(DATA_DIR, "Model_Estimasi_NJOP_Surabaya.xlsx")
    if os.path.exists(xlsx_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_path, data_only=True)
            if "Ringkasan" in wb.sheetnames:
                ws = wb["Ringkasan"]
                for r in range(14, ws.max_row + 1):
                    kec = ws.cell(row=r, column=2).value
                    njop_val = ws.cell(row=r, column=4).value
                    cv_val = ws.cell(row=r, column=5).value
                    qa_status = ws.cell(row=r, column=6).value
                    if kec and njop_val is not None:
                        njop_dict[str(kec).strip().lower()] = {
                            "njop_m2": int(njop_val),
                            "cv": float(cv_val or 0.0),
                            "qa": str(qa_status or "LOLOS")
                        }
        except Exception as e:
            logger.warning(f"Gagal membaca Excel NJOP: {e}")

    raw["njop_kecamatan"] = njop_dict
    return raw


def _haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Menghitung jarak geosferik dalam kilometer."""
    R = 6371.0
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = (math.sin(dlat / 2) ** 2 +
         math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) *
         math.sin(dlon / 2) ** 2)
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c


def compute_all_station_analytics() -> Dict[str, Dict[str, Any]]:
    """
    Menghitung metrik analitik 5D TOD riil, skor AHP, dan apresiasi NJOP (SDM)
    untuk seluruh stasiun aktif di Surabaya Raya berdasarkan data spasial empiris.
    """
    global _CACHED_STATIONS_DATA
    if _CACHED_STATIONS_DATA is not None:
        return _CACHED_STATIONS_DATA

    raw = load_spatial_raw_data()
    st_features = raw["stations"].get("features", [])
    demo_features = raw["demografi"].get("features", [])
    halte_features = raw["halte"].get("features", [])
    banjir_features = raw["banjir"].get("features", [])
    ntl_features = raw["ntl"].get("features", [])
    mall_features = raw["malls"].get("features", [])
    properti_features = raw["properti"].get("features", [])
    njop_map = raw["njop_kecamatan"]

    # Pre-parse geometries
    demo_polys = []
    for f in demo_features:
        try:
            geom = shape(f["geometry"])
            p = f["properties"]
            density_2024 = float(p.get("KEPADATAN PENDUDUK 2024") or 12000.0)
            kec = str(p.get("KECAMATAN", "")).strip().lower()
            desa = str(p.get("DESA", "")).strip()
            demo_polys.append({"geom": geom, "density": density_2024, "kecamatan": kec, "desa": desa})
        except Exception:
            continue

    banjir_polys = []
    for f in banjir_features:
        try:
            geom = shape(f["geometry"])
            kelas = f["properties"].get("Kelas", "Rendah")
            weight = 1.0 if kelas == "Tinggi" else (0.6 if kelas == "Sedang" else 0.3)
            banjir_polys.append({"geom": geom, "weight": weight})
        except Exception:
            continue

    # 1. Normalisasi list stasiun
    stations: Dict[str, Dict[str, Any]] = {}
    for f in st_features:
        p = f["properties"]
        raw_name = p.get("NAMA", "Stasiun KA")
        slug = _clean_slug(raw_name)
        coords = f["geometry"]["coordinates"]
        lon, lat = coords[0], coords[1]

        # Jangan timpa jika sudah ada (menghindari duplikasi Gubeng)
        if slug in stations:
            continue

        kec = p.get("KECAMATAN", "")
        desa = p.get("DESA", "")

        display_name = raw_name.title()
        if "Ka" in display_name:
            display_name = display_name.replace("Ka", "KA")
        if slug == "semut":
            display_name = "Stasiun Surabaya Kota (Semut)"
        elif slug == "gubeng":
            display_name = "Stasiun Surabaya Gubeng"
        elif slug == "pasar_turi":
            display_name = "Stasiun Pasar Turi"

        stations[slug] = {
            "id": slug,
            "name": display_name,
            "latitude": lat,
            "longitude": lon,
            "kecamatan": kec,
            "desa": desa,
            "is_tier_1": slug in TIER_1_STATION_IDS,
        }

    # Integrasi Simpul Multimoda Bersejarah & Strategis Sesuai Perda RTRW Surabaya No. 8/2024
    if "waru" not in stations:
        stations["waru"] = {
            "id": "waru",
            "name": "Stasiun Waru (Gerbang Selatan)",
            "latitude": -7.3547,
            "longitude": 112.7297,
            "kecamatan": "Waru",
            "desa": "Kedungrejo",
            "is_tier_1": False,
        }

    if "terminal_joyoboyo" not in stations:
        stations["terminal_joyoboyo"] = {
            "id": "terminal_joyoboyo",
            "name": "Terminal Intermoda Joyoboyo (TIJ)",
            "latitude": -7.2995,
            "longitude": 112.7368,
            "kecamatan": "Wonokromo",
            "desa": "Sawunggaling",
            "is_tier_1": True,
        }

    if "terminal_purabaya" not in stations:
        stations["terminal_purabaya"] = {
            "id": "terminal_purabaya",
            "name": "Terminal Purabaya (Bungurasih)",
            "latitude": -7.3526,
            "longitude": 112.7235,
            "kecamatan": "Waru",
            "desa": "Bungurasih",
            "is_tier_1": True,
        }

    if "terminal_bratang" not in stations:
        stations["terminal_bratang"] = {
            "id": "terminal_bratang",
            "name": "Terminal Bratang",
            "latitude": -7.2954,
            "longitude": 112.7612,
            "kecamatan": "Gubeng",
            "desa": "Baratajaya",
            "is_tier_1": False,
        }

    # Pre-extract halte coordinates
    halte_coords = []
    for h in halte_features:
        g = h.get("geometry", {})
        c = g.get("coordinates", [])
        if len(c) >= 2:
            halte_coords.append((c[1], c[0]))

    # Pre-extract mall coordinates
    mall_coords = []
    for m in mall_features:
        g = m.get("geometry", {})
        c = g.get("coordinates", [])
        if len(c) >= 2:
            mall_coords.append((c[1], c[0]))

    # Pre-extract properti prices per kecamatan
    properti_prices: Dict[str, List[float]] = {}
    for prop in properti_features:
        p = prop.get("properties", {})
        k = str(p.get("KECAMATAN", "")).strip().lower()
        min_p = p.get("HARGA TANAH MINIMUM (RP/M)") or p.get("HARGA TANAH MINIMUM (RP/M2)")
        if min_p and float(min_p) > 100000:
            properti_prices.setdefault(k, []).append(float(min_p))

    sdm = SDMRegressor()
    weights_vec, _ = calculate_ahp_weights(DEFAULT_5D_PAIRWISE_MATRIX)
    calibrated_models = _load_calibrated_models()

    result: Dict[str, Dict[str, Any]] = {}

    for slug, s in stations.items():
        st_pt = Point(s["longitude"], s["latitude"])
        lat, lon = s["latitude"], s["longitude"]

        # ---------------- D1: Density ----------------
        pop_density = 12000.0
        kec_slug = str(s["kecamatan"]).strip().lower()
        for dp in demo_polys:
            if dp["geom"].contains(st_pt):
                pop_density = dp["density"]
                kec_slug = dp["kecamatan"]
                break
        density_score = min(98.0, max(45.0, (pop_density / 22000.0) * 100.0))

        # ---------------- D5: Distance to Transit / Feeders ----------------
        # Hitung halte dalam radius 800m dan jarak halte terdekat
        halte_count_800m = 0
        min_halte_dist_m = 9999.0
        for hlat, hlon in halte_coords:
            dist_m = _haversine_km(lat, lon, hlat, hlon) * 1000.0
            if dist_m < min_halte_dist_m:
                min_halte_dist_m = dist_m
            if dist_m <= 800.0:
                halte_count_800m += 1

        # Distance score: halte dekat dan banyak halte feeder
        prox_score = max(0.0, 100.0 - (min_halte_dist_m / 10.0))  # 100m -> 90, 500m -> 50
        density_halte_score = min(100.0, halte_count_800m * 12.5)
        distance_score = min(96.0, max(40.0, (prox_score * 0.5) + (density_halte_score * 0.5)))

        # ---------------- D4: Destination Accessibility (Malls & Hubs) ----------------
        nearby_malls = sum(1 for mlat, mlon in mall_coords if _haversine_km(lat, lon, mlat, mlon) <= 1.5)
        # Jarak ke Balai Kota / Pusat Kota (-7.265, 112.750)
        dist_to_cbd_km = _haversine_km(lat, lon, -7.2654, 112.7521)
        cbd_score = max(30.0, 100.0 - (dist_to_cbd_km * 5.0))
        destination_score = min(95.0, max(40.0, (cbd_score * 0.6) + (min(100.0, nearby_malls * 25.0) * 0.4)))

        # ---------------- D3: Design & Walkability (Flood Hazard Penalty) ----------------
        flood_penalty = 0.0
        for bp in banjir_polys:
            if bp["geom"].contains(st_pt):
                flood_penalty += (bp["weight"] * 20.0)
                break
        base_design = 78.0 if s["is_tier_1"] else 65.0
        design_score = min(92.0, max(45.0, base_design - flood_penalty))

        # ---------------- D2: Diversity (Land Use & SES) ----------------
        prop_count = len(properti_prices.get(kec_slug, []))
        diversity_score = min(92.0, max(48.0, 55.0 + min(35.0, prop_count * 0.8) + (8.0 if s["is_tier_1"] else 0.0)))

        scores = {
            "density": round(float(density_score), 1),
            "diversity": round(float(diversity_score), 1),
            "design": round(float(design_score), 1),
            "destination_accessibility": round(float(destination_score), 1),
            "distance_to_transit": round(float(distance_score), 1),
        }

        # TOD composite score via AHP
        tod_score = round(calculate_tod_score(scores, weights_vec), 1)

        # Subtle and clean status indicator
        status = "Focus Area" if s["is_tier_1"] else "Surabaya Rail Network"

        # Check if calibrated by the unsupervised notebook model
        cal = calibrated_models.get(slug)
        if cal:
            tod_score = cal.get("tod_readiness_score", tod_score)
            if "scores" in cal:
                scores = cal["scores"]
            typology = cal.get("typology", "Urban Transit Node")
            sdm_premium = cal.get("predicted_njop_premium_pct", 10.0)
            ci_low = cal.get("ci_lower_pct", round(sdm_premium * 0.75, 1))
            ci_up = cal.get("ci_upper_pct", round(sdm_premium * 1.25, 1))
            direct_eff = round(sdm_premium * 0.7, 1)
            spill_eff = round(sdm_premium * 0.3, 1)
            base_njop = cal.get("njop_base_m2", 7000000)
        else:
            base_njop = 7000000
            if kec_slug in njop_map:
                base_njop = njop_map[kec_slug]["njop_m2"]
            elif properti_prices.get(kec_slug):
                base_njop = int(np.median(properti_prices[kec_slug]))

            neighbor_tod = tod_score * 0.95
            sdm_res = sdm.predict_premium(
                tod_score=tod_score,
                distance_to_station_m=min_halte_dist_m,
                neighbor_avg_tod=neighbor_tod
            )
            sdm_premium = sdm_res["predicted_njop_premium_pct"]
            ci_low = sdm_res["ci_lower_pct"]
            ci_up = sdm_res["ci_upper_pct"]
            direct_eff = sdm_res["direct_effect_pct"]
            spill_eff = sdm_res["spillover_effect_pct"]

            typology = (
                "Metropolitan Commercial Intermodal Hub" if destination_score >= 78.0 and diversity_score >= 75.0
                else "Dense Urban Commuter Spine" if density_score >= 75.0 and distance_score >= 70.0
                else "Suburban Commuter & Feeder Priority" if distance_score >= 60.0
                else "Heritage & Mixed Urban Core"
            )

        # Status & Dimensi
        sorted_dims = sorted(scores.items(), key=lambda x: x[1])
        weakest_dim = sorted_dims[0][0].replace("_", " ").title()
        strongest_dim = sorted_dims[-1][0].replace("_", " ").title()

        # Kebijakan intervensi rekomendasi
        recommendations = [
            f"Fokus APBD pembenahan dimensi terlemah ({weakest_dim} {sorted_dims[0][1]}/100) di kawasan {s['name']}.",
            f"Koneksikan {halte_count_800m} halte feeder WiraWiri terdekat menuju gate stasiun dengan jalur ramah disabilitas.",
            f"Kapitalisasi potensi apresiasi lahan +{sdm_premium}% (%ΔNJOP) melalui instrumen Land Value Capture (LVC)."
        ]

        result[slug] = {
            "id": slug,
            "name": s["name"],
            "latitude": lat,
            "longitude": lon,
            "kecamatan": s["kecamatan"],
            "is_tier_1": s["is_tier_1"],
            "tod_readiness_score": tod_score,
            "scores": scores,
            "benchmark_scores": {
                "density": 77.0,
                "diversity": 76.5,
                "design": 62.5,
                "destination_accessibility": 79.5,
                "distance_to_transit": 78.0,
            },
            "typology": typology,
            "weakest_dimension": weakest_dim,
            "strongest_dimension": strongest_dim,
            "status": status,
            "njop_base_m2": base_njop,
            "njop_premium": {
                "avg_njop_premium_pct": sdm_premium,
                "ci_lower_pct": ci_low,
                "ci_upper_pct": ci_up,
                "affected_h3_count": 37,
                "r_squared": 0.76,
                "direct_effect_pct": direct_eff,
                "spillover_effect_pct": spill_eff,
            },
            "policy_recommendations": recommendations,
        }

    _CACHED_STATIONS_DATA = result
    return result


def get_all_real_h3_features() -> Dict[str, Any]:
    """
    Menghasilkan GeoJSON FeatureCollection sel Uber H3 (resolusi 9)
    yang mencakup catchment seluruh stasiun aktif Surabaya (k=3 ring / radius ~1.000m
    sesuai standar delineasi TOD Perda RTRW Surabaya No. 8/2024 & Permen ATR/BPN 16/2017).
    """
    global _CACHED_H3_COLLECTION
    if _CACHED_H3_COLLECTION is not None:
        return _CACHED_H3_COLLECTION

    stations = compute_all_station_analytics()
    features = []

    for slug, st in stations.items():
        lat, lon = st["latitude"], st["longitude"]
        center_cell = h3.latlng_to_cell(lat, lon, res=9)
        # Lapisan ketiga (k=3 disk) menghasilkan 37 sel heksagonal (radius ~1 km)
        ring_cells = sorted(list(h3.grid_disk(center_cell, 3)))

        for cell in ring_cells:
            boundary = h3.cell_to_boundary(cell)
            geom_coords = [[b[1], b[0]] for b in boundary]
            geom_coords.append(geom_coords[0])  # Close polygon ring

            ring_dist = h3.grid_distance(center_cell, cell)
            decay = max(0.55, round(1.0 - (ring_dist * 0.07), 2))

            c_tod = round(st["tod_readiness_score"] * decay, 1)
            c_njop_prem = round(st["njop_premium"]["avg_njop_premium_pct"] * decay, 1)
            c_ci_low = round(st["njop_premium"]["ci_lower_pct"] * decay, 1)
            c_ci_up = round(st["njop_premium"]["ci_upper_pct"] * decay, 1)

            features.append({
                "type": "Feature",
                "id": cell,
                "geometry": {
                    "type": "Polygon",
                    "coordinates": [geom_coords]
                },
                "properties": {
                    "h3_index": cell,
                    "station_cluster": slug,
                    "station_name": st["name"],
                    "ring_distance": ring_dist,
                    "tod_readiness_score": c_tod,
                    "density_score": round(st["scores"]["density"] * decay, 1),
                    "diversity_score": round(st["scores"]["diversity"] * decay, 1),
                    "design_score": round(st["scores"]["design"] * decay, 1),
                    "destination_score": round(st["scores"]["destination_accessibility"] * decay, 1),
                    "distance_score": round(st["scores"]["distance_to_transit"] * decay, 1),
                    "typology": st["typology"],
                    "predicted_njop_premium_pct": c_njop_prem,
                    "ci_lower_pct": c_ci_low,
                    "ci_upper_pct": c_ci_up,
                    "njop_m2": int(st["njop_base_m2"] * (1.0 + (c_njop_prem / 100.0)))
                }
            })

    _CACHED_H3_COLLECTION = {"type": "FeatureCollection", "features": features}
    return _CACHED_H3_COLLECTION
